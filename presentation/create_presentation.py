"""Generate the T-Analyst presentation deck."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "presentation" / "T-Analyst-Presentation.pptx"
SUMMARY = ROOT / "artifacts" / "recent_summary.json"

W = Inches(13.333)
H = Inches(7.5)

NAVY = RGBColor(16, 42, 58)
DARK = RGBColor(7, 27, 40)
TEAL = RGBColor(0, 122, 115)
TEAL_LIGHT = RGBColor(230, 243, 241)
WHITE = RGBColor(255, 255, 255)
PAPER = RGBColor(247, 249, 250)
MUTED = RGBColor(91, 108, 118)
LINE = RGBColor(219, 228, 232)
GREEN = RGBColor(0, 132, 61)
RED = RGBColor(218, 41, 28)
ORANGE = RGBColor(237, 139, 0)
BLUE = RGBColor(0, 61, 165)
WARM = RGBColor(255, 244, 232)


def set_background(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    color=NAVY,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font="Arial",
    margin=0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_text(slide, runs, x, y, w, h, size=18, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for text, bold, run_color in runs:
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = run_color or color
    return box


def add_box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_line(slide, x1, y1, x2, y2, color=LINE, width=1.5):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_header(slide, kicker, title, subtitle=None):
    add_text(slide, kicker.upper(), 0.65, 0.35, 5.8, 0.3, 10, TEAL, True)
    add_text(slide, title, 0.65, 0.72, 12.0, 0.62, 28, NAVY, True)
    if subtitle:
        add_text(slide, subtitle, 0.65, 1.37, 12.0, 0.45, 15, MUTED)


def add_footer(slide, number):
    add_line(slide, 0.65, 7.1, 12.68, 7.1, LINE, 0.8)
    add_text(
        slide,
        "T-Analyst · Agentic AI",
        0.65,
        7.16,
        4,
        0.18,
        8,
        MUTED,
    )
    add_text(
        slide,
        str(number),
        12.2,
        7.14,
        0.45,
        0.2,
        8,
        MUTED,
        align=PP_ALIGN.RIGHT,
    )


def add_bullets(slide, items, x, y, w, h, size=17, color=NAVY, spacing=9):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.level = 0
        paragraph.text = "•  " + paragraph.text
        paragraph.space_after = Pt(spacing)
    return box


def add_step(slide, number, title, body, x, y, w):
    add_box(slide, x, y, w, 1.35, WHITE, LINE)
    add_text(slide, number, x + 0.22, y + 0.2, 0.45, 0.3, 11, TEAL, True)
    add_text(slide, title, x + 0.22, y + 0.48, w - 0.44, 0.3, 17, NAVY, True)
    add_text(slide, body, x + 0.22, y + 0.83, w - 0.44, 0.37, 11, MUTED)


def slide_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, DARK)
    # MBTA-inspired route lines
    for y, color, width in [
        (0.55, RED, 1.9),
        (0.79, ORANGE, 2.7),
        (1.03, BLUE, 3.5),
        (1.27, GREEN, 4.1),
    ]:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(9.15),
            Inches(y),
            Inches(width),
            Inches(0.09),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
    add_text(slide, "AGENTIC AI · TECHNICAL EVALUATION", 0.75, 0.65, 6, 0.3, 11, RGBColor(119, 215, 199), True)
    add_text(slide, "T-Analyst", 0.75, 1.55, 7.4, 0.85, 36, WHITE, True)
    add_text(
        slide,
        "An Auditable Agent for\nMBTA Transit Operations",
        0.75,
        2.45,
        8.2,
        1.3,
        27,
        WHITE,
        True,
    )
    add_text(
        slide,
        "Ask transit operations questions in plain language—and inspect the evidence behind every number.",
        0.75,
        4.05,
        7.4,
        0.85,
        17,
        RGBColor(209, 226, 232),
    )
    for label, x, w in [
        ("LIVE MBTA DATA", 0.75, 1.75),
        ("REALIZED OPERATIONS", 2.7, 2.2),
        ("EVIDENCE-FIRST AI", 5.1, 2.0),
    ]:
        add_box(slide, x, 5.25, w, 0.48, DARK, RGBColor(66, 93, 105))
        add_text(slide, label, x, 5.38, w, 0.18, 9, WHITE, True, PP_ALIGN.CENTER)
    add_text(
        slide,
        "github.com/OliviaLJY/T-Analyst-An-Auditable-Agent-for-MBTA-Transit-Operations",
        0.75,
        6.75,
        11.8,
        0.25,
        9,
        RGBColor(160, 185, 195),
    )


def slide_problem(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(
        slide,
        "Problem",
        "Public transit data is open. Useful answers are still hard.",
        "The challenge is not access alone—it is interpretation and trust.",
    )
    add_box(slide, 0.65, 2.05, 4.25, 3.95, NAVY, NAVY)
    add_text(slide, "A simple question", 1.0, 2.45, 2.5, 0.3, 12, RGBColor(119, 215, 199), True)
    add_text(
        slide,
        "“Which line had the most uneven train spacing this week?”",
        1.0,
        3.0,
        3.55,
        1.4,
        25,
        WHITE,
        True,
    )
    add_text(
        slide,
        "Answering it requires much more than an API call.",
        1.0,
        5.15,
        3.45,
        0.45,
        13,
        RGBColor(209, 226, 232),
    )
    items = [
        ("01", "Fragmented sources", "Live feeds and completed performance data answer different questions."),
        ("02", "Transit semantics", "A prediction gap is not the same thing as a realized headway."),
        ("03", "Trust gap", "A fluent answer is not useful if the underlying number cannot be checked."),
    ]
    for idx, (n, title, body) in enumerate(items):
        y = 2.05 + idx * 1.32
        add_text(slide, n, 5.45, y + 0.08, 0.5, 0.3, 11, TEAL, True)
        add_text(slide, title, 6.05, y, 2.8, 0.35, 18, NAVY, True)
        add_text(slide, body, 6.05, y + 0.42, 5.85, 0.55, 14, MUTED)
        if idx < 2:
            add_line(slide, 6.05, y + 1.13, 12.2, y + 1.13)
    add_text(
        slide,
        "Design goal: make transit analysis approachable without hiding the method.",
        5.45,
        6.15,
        6.8,
        0.45,
        16,
        TEAL,
        True,
    )
    add_footer(slide, number)


def slide_product(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(
        slide,
        "Product",
        "One interface, two levels of detail",
        "A clear answer first; the technical trail remains one click away.",
    )
    add_box(slide, 0.65, 2.0, 5.75, 4.55, WHITE, LINE)
    add_text(slide, "OVERVIEW", 0.95, 2.25, 1.2, 0.2, 10, TEAL, True)
    add_text(slide, "How is the subway running?", 0.95, 2.62, 4.6, 0.4, 21, NAVY, True)
    for idx, (line_name, value, color) in enumerate(
        [("Red", "15", RED), ("Orange", "8", ORANGE), ("Blue", "7", BLUE), ("Green", "43", GREEN)]
    ):
        x = 0.95 + idx * 1.27
        add_box(slide, x, 3.25, 1.08, 1.12, WHITE, LINE)
        marker = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(3.25), Inches(1.08), Inches(0.07)
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = color
        marker.line.fill.background()
        add_text(slide, line_name, x + 0.12, 3.48, 0.85, 0.2, 10, MUTED, True)
        add_text(slide, value, x + 0.12, 3.77, 0.8, 0.35, 20, NAVY, True)
    add_box(slide, 0.95, 4.72, 4.83, 1.25, TEAL_LIGHT, TEAL_LIGHT)
    add_text(slide, "What stands out", 1.17, 4.96, 2.0, 0.25, 11, TEAL, True)
    add_text(
        slide,
        "Green-D had the largest share of unusually long train intervals: 17.5%.",
        1.17,
        5.3,
        4.25,
        0.48,
        13,
        NAVY,
    )
    add_box(slide, 6.7, 2.0, 5.98, 4.55, WHITE, LINE)
    add_text(slide, "ASK A QUESTION", 7.0, 2.25, 1.8, 0.2, 10, TEAL, True)
    add_text(slide, "What would you like to know?", 7.0, 2.62, 4.9, 0.4, 21, NAVY, True)
    add_box(slide, 7.0, 3.25, 5.05, 0.78, PAPER, LINE)
    add_text(
        slide,
        "How has the Red Line been running?",
        7.22,
        3.5,
        4.5,
        0.25,
        14,
        NAVY,
    )
    add_box(slide, 7.0, 4.35, 5.05, 1.28, TEAL_LIGHT, TEAL_LIGHT)
    add_text(slide, "Answer", 7.22, 4.58, 0.8, 0.2, 10, TEAL, True)
    add_text(
        slide,
        "Recent Red Line spacing was close to schedule at the median, while 9.3% of usable intervals were unusually long. [E1]",
        7.22,
        4.9,
        4.5,
        0.55,
        13,
        NAVY,
    )
    add_text(
        slide,
        "See how this answer was made",
        7.0,
        5.95,
        3.4,
        0.25,
        12,
        TEAL,
        True,
    )
    add_footer(slide, number)


def slide_data(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(
        slide,
        "Data",
        "Live context + realized operations",
        "The app keeps future predictions separate from completed service.",
    )
    add_box(slide, 0.65, 2.0, 5.82, 3.45, WHITE, LINE)
    add_text(slide, "LIVE · MBTA V3 API", 0.95, 2.3, 2.7, 0.25, 11, BLUE, True)
    add_text(slide, "What is happening now?", 0.95, 2.72, 4.2, 0.35, 21, NAVY, True)
    add_bullets(
        slide,
        [
            "/vehicles · reported positions",
            "/alerts · current service notices",
            "/predictions · future arrivals/departures",
        ],
        0.95,
        3.33,
        4.7,
        1.45,
        15,
    )
    add_text(
        slide,
        "Refreshed as a timestamped snapshot",
        0.95,
        5.0,
        4.8,
        0.25,
        11,
        MUTED,
    )
    add_box(slide, 6.82, 2.0, 5.86, 3.45, WHITE, LINE)
    add_text(slide, "HISTORY · MBTA LAMP", 7.12, 2.3, 2.7, 0.25, 11, GREEN, True)
    add_text(slide, "How did service actually run?", 7.12, 2.72, 4.6, 0.35, 21, NAVY, True)
    add_bullets(
        slide,
        [
            "Daily subway performance Parquet",
            "Observed + matched scheduled headways",
            "Travel time, dwell time, route, stop, direction",
        ],
        7.12,
        3.33,
        4.8,
        1.45,
        15,
    )
    add_text(
        slide,
        "Latest 7 complete days · cached locally",
        7.12,
        5.0,
        4.8,
        0.25,
        11,
        MUTED,
    )
    add_rich_text(
        slide,
        [
            ("298,851", True, TEAL),
            (" trip-stop records · ", False, MUTED),
            ("Sep 6–12, 2026", True, NAVY),
            (" · Red, Orange, Blue, Green B/C/D/E", False, MUTED),
        ],
        0.85,
        5.95,
        11.7,
        0.45,
        17,
    )
    add_footer(slide, number)


def slide_architecture(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(
        slide,
        "Agent design",
        "The model reasons. Deterministic tools calculate.",
        "This boundary is the central design choice.",
    )
    labels = [
        ("1", "Question", "Plain language"),
        ("2", "Plan", "GPT-5.5 JSON"),
        ("3", "Validate", "Pydantic"),
        ("4", "Execute", "Python tools"),
        ("5", "Evidence", "[E1], [E2]"),
        ("6", "Explain", "GPT-5.5"),
    ]
    start_x = 0.55
    box_w = 1.75
    gap = 0.34
    y = 2.45
    for idx, (n, title, body) in enumerate(labels):
        x = start_x + idx * (box_w + gap)
        fill = TEAL_LIGHT if idx in (1, 5) else WHITE
        add_box(slide, x, y, box_w, 1.55, fill, LINE)
        add_text(slide, n, x + 0.18, y + 0.18, 0.3, 0.22, 10, TEAL, True)
        add_text(slide, title, x + 0.18, y + 0.52, 1.35, 0.28, 16, NAVY, True)
        add_text(slide, body, x + 0.18, y + 0.95, 1.35, 0.28, 11, MUTED)
        if idx < len(labels) - 1:
            add_line(
                slide,
                x + box_w,
                y + 0.77,
                x + box_w + gap - 0.04,
                y + 0.77,
                TEAL,
                1.7,
            )
    add_box(slide, 1.05, 4.7, 11.25, 1.25, NAVY, NAVY)
    add_text(
        slide,
        "The LLM never calculates a displayed transit metric.",
        1.4,
        4.98,
        5.2,
        0.35,
        19,
        WHITE,
        True,
    )
    add_text(
        slide,
        "It can choose only 4 read-only tools · no arbitrary SQL · unsupported citations trigger a safe fallback",
        1.4,
        5.42,
        9.9,
        0.3,
        13,
        RGBColor(209, 226, 232),
    )
    add_footer(slide, number)


def slide_audit(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(
        slide,
        "Auditability",
        "A fluent answer is not enough",
        "Every numerical claim has a path back to current-run evidence.",
    )
    add_box(slide, 0.65, 2.0, 3.15, 4.35, NAVY, NAVY)
    add_text(slide, "QUESTION", 0.95, 2.32, 1.0, 0.2, 10, RGBColor(119, 215, 199), True)
    add_text(
        slide,
        "How has the Red Line been running?",
        0.95,
        2.8,
        2.55,
        1.0,
        23,
        WHITE,
        True,
    )
    add_text(
        slide,
        "Everyday language in",
        0.95,
        5.55,
        2.4,
        0.25,
        12,
        RGBColor(209, 226, 232),
    )
    add_box(slide, 4.15, 2.0, 3.95, 4.35, WHITE, LINE)
    add_text(slide, "VALIDATED TOOL CALL", 4.45, 2.32, 2.2, 0.2, 10, TEAL, True)
    add_text(slide, "network_reliability", 4.45, 2.8, 3.1, 0.35, 18, NAVY, True)
    add_box(slide, 4.45, 3.35, 3.25, 0.72, PAPER, LINE)
    add_text(slide, '{"line": "Red"}', 4.68, 3.6, 2.7, 0.22, 14, NAVY)
    add_text(slide, "Source", 4.45, 4.48, 1.0, 0.2, 10, MUTED, True)
    add_text(slide, "MBTA LAMP daily subway performance", 4.45, 4.8, 3.1, 0.55, 14, NAVY)
    add_text(slide, "Evidence ID", 4.45, 5.65, 1.1, 0.2, 10, MUTED, True)
    add_text(slide, "[E1]", 5.55, 5.59, 0.75, 0.3, 18, TEAL, True)
    add_box(slide, 8.45, 2.0, 4.23, 4.35, TEAL_LIGHT, TEAL_LIGHT)
    add_text(slide, "GROUNDED ANSWER", 8.75, 2.32, 1.8, 0.2, 10, TEAL, True)
    add_text(
        slide,
        "Recent Red Line spacing was close to schedule at the median.",
        8.75,
        2.82,
        3.5,
        1.0,
        20,
        NAVY,
        True,
    )
    add_text(
        slide,
        "9.3% of usable intervals were more than 1.5× the matched schedule. [E1]",
        8.75,
        4.2,
        3.45,
        0.92,
        16,
        NAVY,
    )
    add_text(
        slide,
        "Answer → evidence out",
        8.75,
        5.55,
        2.8,
        0.25,
        12,
        TEAL,
        True,
    )
    add_footer(slide, number)


def slide_finding(prs, number, summary):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(
        slide,
        "Example finding",
        "Green Line branches showed the most uneven spacing",
        f"Realized headways · {summary['service_dates'][0]} to {summary['service_dates'][-1]}",
    )
    add_box(slide, 0.65, 2.0, 4.0, 4.45, NAVY, NAVY)
    add_text(slide, "HIGHEST LONG-GAP SHARE", 0.98, 2.35, 2.5, 0.2, 10, RGBColor(119, 215, 199), True)
    top = summary["routes"][0]
    add_text(slide, top["route_id"], 0.98, 2.85, 2.3, 0.55, 28, WHITE, True)
    add_text(slide, f"{top['gap_rate'] * 100:.1f}%", 0.98, 3.58, 2.5, 0.75, 34, RGBColor(119, 215, 199), True)
    add_text(
        slide,
        f"{top['observations']:,} usable headway observations",
        0.98,
        4.48,
        2.95,
        0.4,
        14,
        RGBColor(209, 226, 232),
    )
    add_text(
        slide,
        "A long gap is more than 1.5× the matched scheduled headway.",
        0.98,
        5.35,
        2.95,
        0.58,
        13,
        WHITE,
    )
    chart_x, chart_y, chart_w = 5.15, 2.25, 6.85
    rows = summary["routes"]
    max_rate = max(row["gap_rate"] for row in rows)
    for idx, row in enumerate(rows):
        y = chart_y + idx * 0.54
        color = {
            "Red": RED,
            "Orange": ORANGE,
            "Blue": BLUE,
        }.get(row["route_id"], GREEN)
        add_text(slide, row["route_id"], chart_x, y + 0.07, 1.0, 0.22, 12, NAVY, True)
        bar_x = chart_x + 1.05
        bar_w = 4.85 * row["gap_rate"] / max_rate
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(bar_x),
            Inches(y),
            Inches(bar_w),
            Inches(0.33),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        add_text(
            slide,
            f"{row['gap_rate'] * 100:.1f}%",
            bar_x + bar_w + 0.1,
            y + 0.06,
            0.75,
            0.22,
            12,
            NAVY,
            True,
        )
    add_text(
        slide,
        "Share of usable headways > 1.5× schedule",
        6.2,
        6.22,
        4.9,
        0.25,
        11,
        MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, number)


def slide_advantage(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(
        slide,
        "Why it stands out",
        "Trust is a product feature, not a footnote",
        "The project combines agent design with transit operations discipline.",
    )
    add_step(
        slide,
        "01",
        "Bounded agent",
        "Four read-only tools, validated arguments, and no arbitrary SQL.",
        0.65,
        2.15,
        3.8,
    )
    add_step(
        slide,
        "02",
        "Transit-aware",
        "Realized headways stay distinct from future prediction gaps.",
        4.77,
        2.15,
        3.8,
    )
    add_step(
        slide,
        "03",
        "Reproducible",
        "Official sources, cached service days, tests, and generated artifacts.",
        8.89,
        2.15,
        3.8,
    )
    add_box(slide, 0.65, 4.05, 12.04, 1.8, WHITE, LINE)
    add_text(slide, "NOT JUST A DASHBOARD", 0.95, 4.35, 2.2, 0.2, 10, TEAL, True)
    add_text(
        slide,
        "The interface can answer a new question and show how it reached the answer.",
        0.95,
        4.75,
        5.25,
        0.75,
        20,
        NAVY,
        True,
    )
    add_line(slide, 6.55, 4.4, 6.55, 5.53, LINE, 1)
    add_text(slide, "NOT JUST A CHATBOT", 6.9, 4.35, 2.2, 0.2, 10, TEAL, True)
    add_text(
        slide,
        "The model cannot invent a metric path: evidence is produced by deterministic code.",
        6.9,
        4.75,
        5.1,
        0.75,
        20,
        NAVY,
        True,
    )
    add_footer(slide, number)


def slide_limits(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(
        slide,
        "Limits and next steps",
        "A deliberately narrow prototype",
        "The current constraints are visible—and create a clear research path.",
    )
    add_text(slide, "CURRENT LIMITS", 0.75, 2.05, 2.0, 0.22, 10, ORANGE, True)
    add_bullets(
        slide,
        [
            "Seven service days are not a long-run baseline",
            "No weekday/weekend or time-of-day controls",
            "Four tools; no trip planning or forecasting",
            "Local prototype rather than a deployed service",
        ],
        0.75,
        2.5,
        5.2,
        2.8,
        17,
        NAVY,
        14,
    )
    add_line(slide, 6.5, 2.05, 6.5, 5.95, LINE, 1)
    add_text(slide, "NEXT", 6.9, 2.05, 2.0, 0.22, 10, TEAL, True)
    add_bullets(
        slide,
        [
            "Separate peak, off-peak, weekday, and weekend service",
            "Compare Green Line branch and shared-trunk headways",
            "Add data-quality confidence to each evidence package",
            "Evaluate tool selection and answer grounding at scale",
        ],
        6.9,
        2.5,
        5.3,
        2.8,
        17,
        NAVY,
        14,
    )
    add_box(slide, 0.75, 5.85, 11.45, 0.65, WARM, WARM)
    add_text(
        slide,
        "The goal was not to answer every transit question. It was to make a useful set of answers trustworthy.",
        1.0,
        6.05,
        10.95,
        0.25,
        14,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    add_footer(slide, number)


def slide_close(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, DARK)
    add_text(slide, "T-ANALYST", 0.8, 0.7, 2.0, 0.25, 11, RGBColor(119, 215, 199), True)
    add_text(
        slide,
        "A useful agent is not the one\nthat answers everything.",
        0.8,
        1.45,
        7.6,
        1.25,
        30,
        WHITE,
        True,
    )
    add_text(
        slide,
        "It is the one that shows why its answer should be trusted.",
        0.8,
        3.0,
        8.3,
        0.75,
        22,
        RGBColor(119, 215, 199),
        True,
    )
    add_box(slide, 0.8, 4.35, 6.9, 1.15, DARK, RGBColor(66, 93, 105))
    add_text(slide, "Suggested live demo", 1.08, 4.62, 2.0, 0.2, 10, RGBColor(160, 185, 195), True)
    add_text(
        slide,
        "“How has the Red Line been running?”",
        1.08,
        4.96,
        5.9,
        0.3,
        17,
        WHITE,
        True,
    )
    add_text(slide, "Questions?", 9.65, 1.7, 2.6, 0.6, 30, WHITE, True, PP_ALIGN.CENTER)
    add_text(
        slide,
        "github.com/OliviaLJY/\nT-Analyst-An-Auditable-Agent-for-\nMBTA-Transit-Operations",
        8.65,
        3.0,
        4.6,
        1.2,
        12,
        RGBColor(209, 226, 232),
        False,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "MIT Parley GPT-5.5 · MBTA V3 API · MBTA LAMP",
        0.8,
        6.75,
        8,
        0.25,
        10,
        RGBColor(160, 185, 195),
    )
    add_text(slide, str(number), 12.2, 6.75, 0.4, 0.25, 9, RGBColor(160, 185, 195), align=PP_ALIGN.RIGHT)


def build():
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = (
        "T-Analyst: An Auditable Agent for MBTA Transit Operations"
    )
    prs.core_properties.subject = "Agentic AI technical evaluation"
    prs.core_properties.author = "OliviaLJY"
    prs.core_properties.comments = (
        "Generated from the project's verified MBTA LAMP summary."
    )

    slide_title(prs)
    slide_problem(prs, 2)
    slide_product(prs, 3)
    slide_data(prs, 4)
    slide_architecture(prs, 5)
    slide_audit(prs, 6)
    slide_finding(prs, 7, summary)
    slide_advantage(prs, 8)
    slide_limits(prs, 9)
    slide_close(prs, 10)

    prs.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()
